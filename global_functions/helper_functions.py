# ======================== Imports ========================

import sys, os
import xml.etree.ElementTree as ET
import pandas as pd
import json
from openpyxl.styles import PatternFill, Font

def resource_path(relative_path):
    base_path = getattr(sys, "_MEIPASS", os.path.dirname(sys.executable))
    return relative_path
    return os.path.join(base_path, relative_path)


DISPLAY_COLUMNS = [
    "Severity",
    "Alarm Time",
    "Alarm Number",
    "Alarm Text",
    "Supplementary Information",
    "Distinguished Name",
    "Diagnostic Info",
    "Name",
    "Count",
    "DN_Base"
]

DISPLAY_COLUMNS_NEW = [
    "Severity",
    "Alarm Time",
    "Cancel Time",
    "Alarm Number",
    "Alarm Text",
    "Supplementary Information",
    "Distinguished Name",
    "History Match",
    "Status",
    "Resolved",
    "Diagnostic Info",
    "Name",
    "History Count",
]

DATE_COLUMNS = {
    "Alarm Time",
    "Cancel Time",
    "Alarm Insertion Time",
    "Alarm Update Time",
    "Origin Alarm Time",
    "Origin Cancel Time",
}

# =========================
# XML PARSER (namespace-safe)
# =========================
def strip_ns(tag):
    return tag.split('}')[-1]

def parse_xml(file_path, allowed_classes=None):
    tree = ET.parse(file_path)
    root = tree.getroot()

    data = {}

    for mo in root.iter():
        if strip_ns(mo.tag) != "managedObject":
            continue

        cls = mo.get("class")

        # ✅ ONLY CHANGE: filter classes
        if allowed_classes and cls not in allowed_classes:
            continue

        dn = mo.get("distName")

        if not cls:
            continue

        if cls not in data:
            data[cls] = []

        params = {}

        for child in mo:
            tag = strip_ns(child.tag)

            if tag == "p":
                params[child.get("name")] = child.text

            elif tag == "list":
                list_name = child.get("name")
                items = []

                for item in child:
                    if strip_ns(item.tag) != "item":
                        continue

                    item_dict = {}
                    for p in item:
                        if strip_ns(p.tag) == "p":
                            item_dict[p.get("name")] = p.text

                    items.append(item_dict)

                params[list_name] = json.dumps(items, sort_keys=True)

        data[cls].append({
            "distName": dn.rstrip("/") if dn else dn,
            "params": params
        })

    return data


# =========================
# COMPARE PRE vs POST
# =========================
def compare(pre_data, post_data):
    result = {}

    all_classes = set(pre_data) | set(post_data)

    for cls in all_classes:
        pre_list = pre_data.get(cls, [])
        post_list = post_data.get(cls, [])

        dn_pre = {x["distName"]: x["params"] for x in pre_list}
        dn_post = {x["distName"]: x["params"] for x in post_list}

        all_dns = set(dn_pre) | set(dn_post)

        rows = []

        for dn in all_dns:
            pre_p = dn_pre.get(dn, {})
            post_p = dn_post.get(dn, {})

            all_keys = set(pre_p) | set(post_p)

            for k in all_keys:
                rows.append({
                    "distName": dn,
                    "parameter": k,
                    "Pre": pre_p.get(k),
                    "Post": post_p.get(k)
                })

        df = pd.DataFrame(rows)
        result[cls] = df

    return result


# =========================
# NORMALIZE (for comparison)
# =========================
def normalize(val):
    if val is None:
        return ""
    try:
        return json.dumps(json.loads(val), sort_keys=True)
    except:
        return str(val).strip()


# =========================
# SUMMARY DASHBOARD
# =========================
def build_summary(result):
    rows = []

    for cls, df in result.items():
        if df.empty:
            total = 0
            changed = 0
        else:
            total = len(df)
            changed = (df.apply(
                lambda x: normalize(x["Pre"]) != normalize(x["Post"]), axis=1
            )).sum()

        pct = (changed / total * 100) if total else 0

        rows.append({
            "Class": cls,
            "Total Params": total,
            "Changed": changed,
            "% Changed": round(pct, 2)
        })

    summary_df = pd.DataFrame(rows)
    summary_df.sort_values(by="Changed", ascending=False, inplace=True)

    return summary_df


# =========================
# EXCEL WRITER
# =========================
def write_to_excel(result, output_file, sheet_order=None):

    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:

        # ========= SUMMARY =========
        summary_df = build_summary(result)
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        ws = writer.book["Summary"]

        # format summary
        for cell in ws[1]:
            cell.font = Font(bold=True)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # auto width summary
        for col in ws.columns:
            max_len = max(len(str(c.value)) if c.value else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = max_len + 2

        # ========= SHEETS =========
        written = set()

        def write_sheet(cls, df):
            print(f"Working with {cls}")
            df.to_excel(writer, sheet_name=cls[:31], index=False)
            ws = writer.book[cls[:31]]

            # Freeze + filter
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions

            # Bold header
            for cell in ws[1]:
                cell.font = Font(bold=True)

            headers = [c.value for c in ws[1]]
            pre_col = headers.index("Pre") + 1
            post_col = headers.index("Post") + 1

            # # Highlight changes
            # for r in range(2, ws.max_row + 1):
            #     pre_cell = ws.cell(r, pre_col)
            #     post_cell = ws.cell(r, post_col)
            #
            #     if normalize(pre_cell.value) != normalize(post_cell.value):
            #         pre_cell.fill = red_fill
            #         post_cell.fill = green_fill
            #
            # # Auto width
            # for col in ws.columns:
            #     max_len = max(len(str(c.value)) if c.value else 0 for c in col)
            #     ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

        # write ordered sheets
        if sheet_order:
            for cls in sheet_order:
                df = result.get(cls)
                if df is not None and not df.empty:
                    write_sheet(cls, df)
                    written.add(cls)

        # write remaining
        for cls, df in result.items():
            if cls in written or df.empty:
                continue
            write_sheet(cls, df)